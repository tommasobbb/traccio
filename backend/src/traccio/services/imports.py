"""Decode an uploaded import file into rows (ADR 0023).

This is the only place in the import path that touches a file format. It turns
raw bytes — an ``.xlsx`` workbook or a CSV — into a list of ``header -> cell``
maps, then hands off to the pure
:func:`traccio.domain.imports.parse.parse_import`. Reading order, sheet
selection, delimiter sniffing and BOM handling live here; classification does
not.

Data safety (``.claude/rules/data-safety.md``): nothing here logs a cell value,
a description, or an amount — only counts and stable reason codes.
"""

import csv
import io

from openpyxl import load_workbook

# ``.xlsx`` is a ZIP; every one starts with this local-file-header signature.
_ZIP_MAGIC = b"PK\x03\x04"
# Guard the CSV sniffer against a pathological first line.
_SNIFF_BYTES = 8192


class ImportDecodeError(Exception):
    """The uploaded bytes are not a readable import file.

    Carries a stable, value-free ``reason`` (``"empty_file"``,
    ``"unreadable_file"``, ``"no_header_row"``) so the API can map it to a
    ``422`` without inspecting the message.
    """

    def __init__(self, reason: str) -> None:
        super().__init__(f"cannot decode import file: {reason}")
        self.reason = reason


def decode_rows(content: bytes, *, filename: str) -> tuple[list[str], list[dict[str, object]]]:
    """Decode ``content`` into its header and one ``header -> cell`` map per row.

    The format is chosen by content sniffing (the ZIP magic) first, then the
    filename extension. Header cells are trimmed; a row shorter than the header
    is padded with ``None``, a longer one is truncated.

    Parameters
    ----------
    content : bytes
        The raw uploaded file.
    filename : str
        The client-supplied name, used only as a fallback format hint.

    Returns
    -------
    tuple[list[str], list[dict[str, object]]]
        The trimmed header cells, and the data rows keyed by header. xlsx cells
        keep their openpyxl type (``str``/``int``/``float``/``datetime``/
        ``None``); CSV cells are ``str``. The header is returned even when
        there are zero data rows, so the caller can still check for missing
        columns.

    Raises
    ------
    ImportDecodeError
        The bytes are empty, unreadable, or carry no header row.
    """
    if not content:
        raise ImportDecodeError("empty_file")

    is_xlsx = content[:4] == _ZIP_MAGIC or filename.lower().endswith(".xlsx")
    header, records = _read_xlsx(content) if is_xlsx else _read_csv(content)

    if not header:
        raise ImportDecodeError("no_header_row")

    rows: list[dict[str, object]] = []
    for record in records:
        padded: list[object] = [*record, *([None] * (len(header) - len(record)))]
        rows.append({key: padded[i] for i, key in enumerate(header)})
    return header, rows


def _read_xlsx(content: bytes) -> tuple[list[str], list[list[object]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of types on a bad file
        raise ImportDecodeError("unreadable_file") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise ImportDecodeError("no_header_row")
        iterator = sheet.iter_rows(values_only=True)
        try:
            first = next(iterator)
        except StopIteration:
            return [], []
        header = [str(cell).strip() for cell in first if cell is not None]
        body: list[list[object]] = [
            list(row) for row in iterator if any(cell is not None for cell in row)
        ]
        return header, body
    finally:
        workbook.close()


def _read_csv(content: bytes) -> tuple[list[str], list[list[object]]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:_SNIFF_BYTES]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        header = [cell.strip() for cell in next(reader)]
    except StopIteration:
        return [], []
    body: list[list[object]] = [list(row) for row in reader if any(cell.strip() for cell in row)]
    return header, body
